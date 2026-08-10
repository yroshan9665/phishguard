from functools import wraps
from io import BytesIO
from datetime import datetime

from flask import Blueprint, render_template, redirect, url_for, flash, jsonify, send_file, request
from flask_login import login_required, current_user, login_user
from sqlalchemy import func
from app import db, bcrypt
from app.models import User, ScanResult, LoginLog

admin_bp = Blueprint('admin', __name__, url_prefix='/admin')


def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated:
            return redirect(url_for('admin.admin_login'))
        if not current_user.is_admin:
            flash('Admin access required.', 'danger')
            return redirect(url_for('dashboard.index'))
        return f(*args, **kwargs)
    return decorated


@admin_bp.route('/login', methods=['GET', 'POST'])
def admin_login():
    # If a regular user is logged in, send them to their dashboard
    if current_user.is_authenticated:
        if current_user.is_admin:
            return redirect(url_for('admin.index'))
        return redirect(url_for('dashboard.index'))
    if request.method == 'POST':
        email    = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        user     = User.query.filter_by(email=email, is_admin=True).first()
        if user and bcrypt.check_password_hash(user.password, password):
            login_user(user)
            db.session.add(LoginLog(user_id=user.id, ip_address=request.remote_addr, status='success'))
            db.session.commit()
            flash(f'Welcome, {user.username}. Admin access granted.', 'success')
            return redirect(url_for('admin.index'))
        # log failed attempt if user exists but wrong role/password
        non_admin = User.query.filter_by(email=email).first()
        if non_admin:
            db.session.add(LoginLog(user_id=non_admin.id, ip_address=request.remote_addr, status='failed'))
            db.session.commit()
        flash('Invalid credentials or not an admin account.', 'danger')
        return redirect(url_for('admin.admin_login'))
    return render_template('admin/login.html')


@admin_bp.route('/')
@admin_required
def index():
    users        = User.query.order_by(User.created_at.desc()).all()
    total_scans  = ScanResult.query.count()
    phishing_count   = ScanResult.query.filter_by(prediction='Phishing').count()
    safe_count       = ScanResult.query.filter_by(prediction='Safe').count()
    suspicious_count = ScanResult.query.filter_by(prediction='Suspicious').count()
    recent_scans = (ScanResult.query
                    .order_by(ScanResult.scanned_at.desc())
                    .limit(20).all())
    login_logs = (LoginLog.query
                  .order_by(LoginLog.logged_at.desc())
                  .limit(30).all())
    stats = {
        'total_users':      len(users),
        'total_scans':      total_scans,
        'phishing_count':   phishing_count,
        'safe_count':       safe_count,
        'suspicious_count': suspicious_count,
    }
    return render_template('admin/index.html', users=users,
                           recent_scans=recent_scans, login_logs=login_logs, stats=stats)


@admin_bp.route('/export/excel')
@admin_required
def export_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        flash('openpyxl not installed. Run: pip install openpyxl', 'danger')
        return redirect(url_for('admin.index'))

    wb = openpyxl.Workbook()
    PURPLE = "7C6FFF"; PURPLE_DK = "4A3FCC"; BG_DARK = "0D0D2B"
    BG_ROW1 = "12122E"; BG_ROW2 = "0F0F26"; WHITE = "FFFFFF"
    MUTED = "8888AA"; GREEN = "10B981"; YELLOW = "F59E0B"; RED = "EF4444"

    def hdr_fill(c): return PatternFill("solid", fgColor=c)
    def thin_border():
        s = Side(style='thin', color="1E1E3F")
        return Border(left=s, right=s, top=s, bottom=s)
    def col_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 1 — Users
    ws1 = wb.active; ws1.title = "Users Summary"
    ws1.sheet_view.showGridLines = False
    ws1.merge_cells("A1:H1")
    tc = ws1["A1"]; tc.value = "PhishGuard AI — User Summary Report"
    tc.font = Font(name="Calibri", bold=True, size=16, color=WHITE)
    tc.fill = hdr_fill(PURPLE_DK)
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:H2")
    sc = ws1["A2"]; sc.value = f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    sc.font = Font(name="Calibri", size=10, color=MUTED)
    sc.fill = hdr_fill(BG_DARK)
    sc.alignment = Alignment(horizontal="center", vertical="center")

    headers1 = ["#", "Username", "Email", "Role", "Joined Date", "Total Scans", "Phishing", "Safe"]
    for c, h in enumerate(headers1, 1):
        cell = ws1.cell(4, c, h)
        cell.font = Font(name="Calibri", bold=True, size=11, color=WHITE)
        cell.fill = hdr_fill(PURPLE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border()
    ws1.row_dimensions[4].height = 26

    users = User.query.order_by(User.created_at.desc()).all()
    for r, u in enumerate(users, 5):
        total = ScanResult.query.filter_by(user_id=u.id).count()
        phish = ScanResult.query.filter_by(user_id=u.id, prediction='Phishing').count()
        safe  = ScanResult.query.filter_by(user_id=u.id, prediction='Safe').count()
        role  = "Admin" if u.is_admin else "User"
        row_bg = BG_ROW1 if r % 2 == 0 else BG_ROW2
        vals = [r-4, u.username, u.email, role,
                u.created_at.strftime('%Y-%m-%d'), total, phish, safe]
        for c, v in enumerate(vals, 1):
            cell = ws1.cell(r, c, v)
            cell.font = Font(name="Calibri", size=10, color=WHITE)
            cell.fill = hdr_fill(row_bg)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border()
        ws1.row_dimensions[r].height = 22
    col_widths(ws1, [6, 18, 30, 10, 16, 14, 12, 10])

    # Sheet 2 — Scans
    ws2 = wb.create_sheet("Scan Records")
    ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:I1")
    tc2 = ws2["A1"]; tc2.value = "PhishGuard AI — Complete Scan Records"
    tc2.font = Font(name="Calibri", bold=True, size=16, color=WHITE)
    tc2.fill = hdr_fill("054F63")
    tc2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 36

    headers2 = ["#", "Username", "Email", "URL", "Verdict", "Confidence %", "Risk Score", "Scan Date", "IP Used"]
    for c, h in enumerate(headers2, 1):
        cell = ws2.cell(4, c, h)
        cell.font = Font(name="Calibri", bold=True, size=11, color=WHITE)
        cell.fill = hdr_fill("054F63")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border()
    ws2.row_dimensions[4].height = 26

    scans = ScanResult.query.order_by(ScanResult.scanned_at.desc()).all()
    v_colors = {'Phishing': RED, 'Suspicious': YELLOW, 'Safe': GREEN}
    for r, s in enumerate(scans, 5):
        row_bg = BG_ROW1 if r % 2 == 0 else BG_ROW2
        has_ip = "Yes" if (s.features or {}).get('has_ip', 0) == 1 else "No"
        vals = [r-4, s.user.username, s.user.email, s.url, s.prediction,
                round(s.confidence, 1), s.risk_score,
                s.scanned_at.strftime('%Y-%m-%d %H:%M'), has_ip]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(r, c, v)
            cell.font = Font(name="Calibri", size=10,
                             color=v_colors.get(s.prediction, WHITE) if c == 5 else WHITE)
            cell.fill = hdr_fill(row_bg)
            cell.alignment = Alignment(horizontal="center" if c != 4 else "left", vertical="center")
            cell.border = thin_border()
        ws2.row_dimensions[r].height = 22
    col_widths(ws2, [6, 16, 28, 50, 14, 14, 12, 20, 10])

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    filename = f"phishguard_users_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@admin_bp.route('/user/<int:user_id>/toggle-admin', methods=['POST'])
@admin_required
def toggle_admin(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash('Cannot modify your own admin status.', 'warning')
    else:
        user.is_admin = not user.is_admin
        db.session.commit()
        flash(f'Admin status updated for {user.username}.', 'success')
    return redirect(url_for('admin.index'))


@admin_bp.route('/user/<int:user_id>/delete', methods=['POST'])
@admin_required
def delete_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash('Cannot delete your own account from admin panel.', 'warning')
    else:
        ScanResult.query.filter_by(user_id=user_id).delete()
        db.session.delete(user)
        db.session.commit()
        flash(f'User {user.username} deleted.', 'success')
    return redirect(url_for('admin.index'))


@admin_bp.route('/api/chart-data')
@admin_required
def chart_data():
    from datetime import timedelta
    dist = (db.session.query(ScanResult.prediction, func.count(ScanResult.id))
            .group_by(ScanResult.prediction).all())
    dist_dict = {row[0]: row[1] for row in dist}

    top_users_q = (db.session.query(ScanResult.user_id, func.count(ScanResult.id).label('cnt'))
                   .group_by(ScanResult.user_id)
                   .order_by(func.count(ScanResult.id).desc())
                   .limit(7).all())
    labels, totals, safe_d, susp_d, phish_d = [], [], [], [], []
    for row in top_users_q:
        u = User.query.get(row[0])
        if not u:
            continue
        labels.append(u.username)
        totals.append(row[1])
        safe_d.append(ScanResult.query.filter_by(user_id=row[0], prediction='Safe').count())
        susp_d.append(ScanResult.query.filter_by(user_id=row[0], prediction='Suspicious').count())
        phish_d.append(ScanResult.query.filter_by(user_id=row[0], prediction='Phishing').count())

    # 7-day global scan trend
    today = datetime.utcnow().date()
    trend_labels, trend_data = [], []
    for i in range(6, -1, -1):
        day = today - timedelta(days=i)
        day_start = datetime(day.year, day.month, day.day)
        day_end   = day_start + timedelta(days=1)
        cnt = ScanResult.query.filter(
            ScanResult.scanned_at >= day_start,
            ScanResult.scanned_at < day_end
        ).count()
        trend_labels.append(day.strftime('%a'))
        trend_data.append(cnt)

    return jsonify({
        'distribution': {
            'labels': ['Safe', 'Suspicious', 'Phishing'],
            'data': [dist_dict.get('Safe', 0),
                     dist_dict.get('Suspicious', 0),
                     dist_dict.get('Phishing', 0)],
        },
        'top_users': {
            'labels': labels,
            'data':   totals,
            'safe':   safe_d,
            'suspicious': susp_d,
            'phishing':   phish_d,
        },
        'scan_trend': {'labels': trend_labels, 'data': trend_data},
    })
